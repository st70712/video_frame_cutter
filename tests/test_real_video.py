import json
import os
import time
from itertools import pairwise
from pathlib import Path
from threading import enumerate as active_threads

import pytest
from PIL import Image
from pptx import Presentation
from PySide6.QtCore import QTimer
from PySide6.QtMultimedia import QMediaPlayer

from scripts.benchmark_analysis import AnalysisRecorder, memory_usage, result_signature, save_report
from video_frame_cutter.analysis import analysis_defaults, analyze
from video_frame_cutter.export import export_markers
from video_frame_cutter.models import AnalysisSettings, CropRect, ExportSettings
from video_frame_cutter.project import fingerprint, read_project, save_project, validate_source
from video_frame_cutter.ui.main_window import MainWindow
from video_frame_cutter.video import FrameReader


@pytest.fixture
def monitored_window(qtbot):
    window = MainWindow()
    window.confirm_discard = lambda: True
    qtbot.addWidget(window)
    started = time.monotonic()
    timer = QTimer(window)
    timer.setInterval(5000)
    timer.timeout.connect(lambda: print(
        f"HEARTBEAT elapsed={time.monotonic() - started:.1f}s "
        f"progress={window.progress.value()}% jobs={len(window.jobs)} "
        f"stage={window.statusBar().currentMessage()}", flush=True,
    ))
    timer.start()
    yield window
    timer.stop()
    window.dirty = False
    window.close()
    qtbot.waitUntil(lambda: not window.jobs, timeout=30000)


@pytest.mark.skipif(os.environ.get("VFC_REAL_TEST") != "1", reason="Opt-in original video check")
def test_original_video(qtbot, tmp_path, monitored_window):
    source = Path("video.mp4").resolve()
    print("STAGE source fingerprint", flush=True)
    before = fingerprint(source)
    errors = []
    window = monitored_window
    window.error = errors.append
    window.show()
    started = time.monotonic()
    print("STAGE indexing", flush=True)
    window.load_video(source)
    qtbot.waitUntil(
        lambda: bool(errors) or (window.info is not None and window.frame_image is not None),
        timeout=240000,
    )
    assert not errors
    info = window.info
    print(
        f"INDEX seconds={time.monotonic() - started:.2f} frames={len(info.frames)} duration={info.duration:.3f}"
    )
    assert info.audio
    print("STAGE audio/video playback", flush=True)
    window.toggle_play()
    qtbot.waitUntil(lambda: window.player.position() > 800, timeout=20000)
    assert window.player.playbackState() == QMediaPlayer.PlaybackState.PlayingState
    assert window.video_widget.videoSink().videoFrame().isValid()
    window.volume.setValue(25)
    assert abs(window.audio.volume() - 0.25) < 0.001
    window.mute.setChecked(True)
    assert window.audio.isMuted()
    window.toggle_play()
    qtbot.waitUntil(lambda: window.frame_image is not None)
    print("STAGE manual markers, project roundtrip and exports", flush=True)
    for seconds in (0, info.duration / 2, info.frames[-1].seconds):
        window.seek(seconds)
        qtbot.waitUntil(lambda: window.frame_image is not None)
        window.add_marker()
    markers = window.markers
    assert len(markers) == 3
    markers[1].crop = CropRect(0.1, 0.1, 0.9, 0.85)
    settings = ExportSettings(640, 360, 95, True)
    saved = tmp_path / "real.vfc.json"
    save_project(saved, info, markers, AnalysisSettings(), settings, before)
    document, restored, analysis_settings, export_settings = read_project(saved)
    assert analysis_settings == AnalysisSettings()
    validate_source(document, info, restored, before)
    directory = export_markers(info, restored, export_settings, tmp_path, "jpg")
    output = export_markers(info, restored, export_settings, tmp_path / "real.pptx", "pptx")
    assert len(Presentation(output).slides) == 3
    assert all(Image.open(path).size == (640, 360) for path in directory.glob("*.jpg"))
    with FrameReader(info) as reader:
        assert reader.get(info.frames[-1]).size == (1280, 720)
    window.seek(info.duration / 2)
    qtbot.waitUntil(lambda: window.frame_image is not None)
    verification = Path(os.environ.get("VFC_VERIFY_DIR", "outputs/verification"))
    verification.mkdir(parents=True, exist_ok=True)
    window.resize(1280, 860)
    window.grab().save(str(verification / "desktop.png"))
    window.resize(1024, 720)
    window.grab().save(str(verification / "compact.png"))
    width = int(os.environ.get("VFC_ANALYSIS_WIDTH", "480"))
    repeats = int(os.environ.get("VFC_ANALYSIS_REPEATS", "1"))
    limit = float(os.environ.get("VFC_ANALYSIS_LIMIT", "0"))
    reference_path = os.environ.get("VFC_BASELINE_REPORT")
    expected_signature = None
    if reference_path:
        reference = json.loads(Path(reference_path).read_text(encoding="utf-8"))
        assert reference["source"] == before
        assert reference["settings"]["width"] == width
        assert reference["baseline"]["complete_video"]
        expected_signature = reference["baseline"]["signature"]
    metrics = {"analysis_width": width, "runs": [], "source": before}
    for repeat in range(repeats):
        qtbot.waitUntil(lambda: not window.jobs, timeout=30000)
        pulses = []
        memory_samples = [memory_usage().get("working_set_bytes", 0)]
        timer = QTimer(window)

        def heartbeat(pulses=pulses, memory_samples=memory_samples):
            pulses.append(time.perf_counter())
            memory_samples.append(memory_usage().get("working_set_bytes", 0))

        timer.timeout.connect(heartbeat)
        results = []
        print(f"STAGE full analysis width={width} run={repeat + 1} frames={len(info.frames)}",
              flush=True)
        with AnalysisRecorder() as recorder:
            started = time.perf_counter()
            timer.start(50)
            window.launch_job(
                lambda cancel, progress: analyze(info, AnalysisSettings(width=width), cancel, progress),
                results.append,
                "Test analysis",
            )
            try:
                qtbot.waitUntil(
                    lambda results=results: (bool(results) or bool(errors))
                    and window.busy_job is None,
                    timeout=int(os.environ.get("VFC_TEST_TIMEOUT", "3600")) * 1000,
                )
            finally:
                timer.stop()
                timer.deleteLater()
        ended = time.perf_counter()
        assert not errors
        result = results[0]
        signature = result_signature(result, recorder.pixel_digest.hexdigest())
        timestamps = [started, *pulses, ended]
        measured = {
            "analysis_elapsed_seconds": ended - started,
            "gui_max_gap_seconds": max(
                right - left for left, right in pairwise(timestamps)
            ),
            "gui_pulses": len(pulses),
            "sampled_memory_increase_bytes": max(memory_samples) - memory_samples[0],
            "memory_after": memory_usage(),
            "configuration": analysis_defaults(),
            "signature": signature,
            "matches_baseline": signature == expected_signature if expected_signature else None,
        }
        metrics["runs"].append(measured)
        save_report(verification / "analysis.json", metrics)
        print(f"ANALYSIS seconds={ended - started:.2f} markers={len(result.markers)} "
              f"pulses={len(pulses)} max_gap={measured['gui_max_gap_seconds']:.3f}s", flush=True)
        assert len(pulses) > 10
        assert len(result.curve) == len(info.frames)
        assert window.progress.value() == 100
        assert not any(thread.name.startswith("vfc-analysis") for thread in active_threads())
        if expected_signature is not None:
            assert signature == expected_signature
        if limit:
            assert measured["analysis_elapsed_seconds"] < limit
            assert measured["gui_max_gap_seconds"] < 1
            assert measured["sampled_memory_increase_bytes"] <= 256 * 1024 * 1024

    print("STAGE cancellation", flush=True)
    existing_ids = [marker.uid for marker in window.markers]
    cancelled_results = []
    window.launch_job(
        lambda cancel, progress: analyze(info, AnalysisSettings(width=width), cancel, progress),
        cancelled_results.append, "Cancellation test",
    )
    qtbot.waitUntil(lambda: window.progress.value() >= 1 or bool(errors), timeout=30000)
    assert not errors
    started = time.perf_counter()
    window.cancel_work()
    qtbot.waitUntil(lambda: not window.jobs, timeout=10000)
    metrics["cancel_seconds"] = time.perf_counter() - started
    save_report(verification / "analysis.json", metrics)
    assert not cancelled_results
    assert [marker.uid for marker in window.markers] == existing_ids
    assert not any(thread.name.startswith("vfc-analysis") for thread in active_threads())
    if limit:
        assert metrics["cancel_seconds"] < 2
    assert len(result.curve) == len(info.frames)
    assert result.markers
    stable_markers = [marker for marker in result.markers if marker.included]
    assert stable_markers
    positions = sorted({0, len(stable_markers) // 2, len(stable_markers) - 1})
    sampled = [stable_markers[position] for position in positions]
    print(f"STAGE export analyzed markers sample={len(sampled)} stable={len(stable_markers)}", flush=True)
    directory = export_markers(info, sampled, settings, verification, "jpg")
    output = export_markers(info, sampled, settings, verification / "analyzed.pptx", "pptx")
    deck = Presentation(output)
    files = sorted(directory.glob("*.jpg"))
    assert len(deck.slides) == len(files) == len(sampled)
    for slide, path in zip(deck.slides, files):
        assert slide.shapes[0].image.blob == path.read_bytes()
    assert len({marker.frame.pts for marker in result.markers}) == len(result.markers)
    print("AUDIO backend reports active; subjective audio sync requires human listening")
    print("SOURCE unchanged", before == fingerprint(source))
    assert before == fingerprint(source)
    window.dirty = False
    window.close()
    qtbot.waitUntil(lambda: not window.jobs, timeout=10000)
