import json
from io import BytesIO
from threading import Event

import pytest
from PIL import Image
from pptx import Presentation

from video_frame_cutter.export import export_markers
from video_frame_cutter.models import AnalysisSettings, CropRect, ExportSettings, Marker
from video_frame_cutter.project import fingerprint, read_project, save_project, validate_source
from video_frame_cutter.video import Cancelled, index_video


def test_roundtrip_and_export(video_path, tmp_path):
    info = index_video(video_path)
    markers = [Marker(info.frames[20], crop=CropRect(0.2, 0.1, 0.8, 0.9)), Marker(info.frames[2])]
    settings = ExportSettings(320, 180, 95, True)
    analysis_settings = AnalysisSettings(
        duplicate_window_seconds=3.0, start_seconds=0.5, end_seconds=2.5
    )
    identity = fingerprint(video_path)
    project = tmp_path / "project.json"
    save_project(project, info, markers, analysis_settings, settings, identity)
    document, loaded, analysis, export = read_project(project)
    assert analysis == analysis_settings
    validate_source(document, info, loaded, identity)
    assert loaded == list(reversed(markers))
    assert export == settings
    directory = export_markers(info, markers, settings, tmp_path, "jpg")
    files = sorted(directory.glob("*.jpg"))
    assert len(files) == 2
    for path in files:
        with Image.open(path) as image:
            assert image.size == (320, 180)
            assert image.mode == "RGB"
    output = tmp_path / "output.pptx"
    export_markers(info, markers, settings, output, "pptx")
    deck = Presentation(output)
    assert len(deck.slides) == 2
    for slide, path in zip(deck.slides, files):
        assert slide.shapes[0].image.blob == path.read_bytes()
        assert slide.shapes[1].top >= slide.shapes[0].height
        assert Image.open(BytesIO(slide.shapes[0].image.blob)).size == (320, 180)
    with pytest.raises(FileExistsError):
        export_markers(info, markers, settings, output, "pptx")

    legacy_document = json.loads(project.read_text(encoding="utf-8"))
    for field in ("duplicate_window_seconds", "start_seconds", "end_seconds"):
        legacy_document["analysis"].pop(field)
    legacy_project = tmp_path / "legacy-project.json"
    legacy_project.write_text(json.dumps(legacy_document), encoding="utf-8")
    _, _, legacy_analysis, _ = read_project(legacy_project)
    assert legacy_analysis.duplicate_window_seconds == 0.0
    assert legacy_analysis.start_seconds == 0.0
    assert legacy_analysis.end_seconds is None


def test_cancel_and_source_protection(video_path, tmp_path):
    info = index_video(video_path)
    before = fingerprint(video_path)
    markers = [Marker(info.frames[0])]
    cancel = Event()
    cancel.set()
    with pytest.raises(Cancelled):
        export_markers(info, markers, ExportSettings(), tmp_path, "jpg", cancel)
    assert not list(tmp_path.glob(".frame-export-*"))
    with pytest.raises(ValueError):
        export_markers(info, markers, ExportSettings(), video_path, "pptx", overwrite=True)
    assert fingerprint(video_path) == before


def test_original_fit_keeps_each_crop_at_its_own_pixels(video_path, tmp_path):
    info = index_video(video_path)
    markers = [
        Marker(info.frames[2]),
        Marker(info.frames[20], crop=CropRect(0.2, 0.1, 0.8, 0.9)),
    ]
    settings = ExportSettings(320, 180, 95, True, "original")

    def crop_size(marker):
        left, top, right, bottom = marker.crop.pixels((info.width, info.height))
        return (right - left, bottom - top)

    expected = [crop_size(marker) for marker in markers]
    assert expected[0] != expected[1]
    assert (settings.width, settings.height) not in expected

    directory = export_markers(info, markers, settings, tmp_path, "jpg")
    for path, size in zip(sorted(directory.glob("*.jpg")), expected):
        with Image.open(path) as image:
            assert image.size == size

    output = tmp_path / "original.pptx"
    export_markers(info, markers, settings, output, "pptx")
    deck = Presentation(output)
    pictures = [slide.shapes[0] for slide in deck.slides]
    assert [picture.image.size for picture in pictures] == expected
    assert 914400 <= deck.slide_width <= 51206400
    assert 914400 <= deck.slide_height <= 51206400
    scale = pictures[0].width / expected[0][0]
    for picture, (width, height) in zip(pictures, expected):
        # One shared scale, so no picture is distorted and a smaller crop stays smaller.
        assert picture.width / width == pytest.approx(scale, rel=0.001)
        assert picture.height / height == pytest.approx(scale, rel=0.001)
        assert picture.left >= 0 and picture.top >= 0
        assert picture.left + picture.width <= deck.slide_width
        assert picture.top + picture.height <= deck.slide_height
    footers = [slide.shapes[1] for slide in deck.slides]
    for footer, picture in zip(footers, pictures):
        assert footer.top >= picture.top + picture.height

    stretched = export_markers(
        info, markers, ExportSettings(320, 180, 95, True), tmp_path, "jpg"
    )
    for path in stretched.glob("*.jpg"):
        with Image.open(path) as image:
            assert image.size == (320, 180)


def test_fit_setting_persists_and_is_validated(video_path, tmp_path):
    info = index_video(video_path)
    markers = [Marker(info.frames[0])]
    settings = ExportSettings(320, 180, 95, False, "original")
    project = tmp_path / "fit.json"
    save_project(project, info, markers, AnalysisSettings(), settings, fingerprint(video_path))
    document = json.loads(project.read_text(encoding="utf-8"))
    assert document["export"]["fit"] == "original"
    assert read_project(project)[3] == settings

    legacy = tmp_path / "legacy-fit.json"
    document["export"].pop("fit")
    legacy.write_text(json.dumps(document), encoding="utf-8")
    assert read_project(legacy)[3].fit == "stretch"

    with pytest.raises(ValueError):
        ExportSettings(fit="letterbox").validate()
    broken = tmp_path / "broken-fit.json"
    document["export"]["fit"] = "letterbox"
    broken.write_text(json.dumps(document), encoding="utf-8")
    with pytest.raises(ValueError):
        read_project(broken)


@pytest.mark.parametrize("size", [(8192, 1080), (8192, 16), (16, 8192), (1920, 1080)])
@pytest.mark.parametrize("timestamp", [False, True])
def test_pptx_extreme_ratios(video_path, tmp_path, size, timestamp):
    info = index_video(video_path)
    settings = ExportSettings(*size, timestamp=timestamp)
    output = tmp_path / "ratio.pptx"
    export_markers(info, [Marker(info.frames[0])], settings, output, "pptx")
    deck = Presentation(output)
    picture = deck.slides[0].shapes[0]
    assert 914400 <= deck.slide_width <= 51206400
    assert 914400 <= deck.slide_height <= 51206400
    assert picture.width / picture.height == pytest.approx(size[0] / size[1], rel=0.001)
    assert picture.left >= 0 and picture.top >= 0
    assert picture.left + picture.width <= deck.slide_width
    assert picture.top + picture.height <= deck.slide_height
    assert Image.open(BytesIO(picture.image.blob)).size == size
    if timestamp:
        footer = deck.slides[0].shapes[1]
        assert footer.top >= picture.top + picture.height
