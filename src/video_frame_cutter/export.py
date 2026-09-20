import os
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from .models import timecode
from .video import FrameReader, check_cancel

# Widest and tallest a picture may be placed. PowerPoint rejects slides over 56 inches,
# so the largest picture is kept well inside that and everything else shares its scale.
PICTURE_LIMIT = (Inches(50), Inches(7.5))


def render_marker(reader, marker, settings, cancel=None):
    settings.validate()
    image = reader.get(marker.frame, cancel)
    image = image.crop(marker.crop.pixels(image.size))
    size = settings.output_size(image.size)
    if size != image.size:
        image = image.resize(size, Image.Resampling.LANCZOS)
    output = BytesIO()
    image.convert("RGB").save(output, "JPEG", quality=settings.quality)
    return output.getvalue()


def add_slide(deck, content, marker, timestamp):
    """Add one picture slide, leaving its geometry to :func:`layout_deck`."""
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    picture = slide.shapes.add_picture(BytesIO(content), 0, 0)
    caption = None
    if timestamp:
        caption = slide.shapes.add_textbox(0, 0, 0, 0)
        paragraph = caption.text_frame.paragraphs[0]
        paragraph.text = timecode(marker.frame.seconds)
        paragraph.font.size = Pt(12)
    return picture, caption, picture.image.size


def layout_deck(deck, placements, timestamp):
    """Size the deck for the largest picture and centre the content of every slide.

    A presentation has one slide size for all of its slides, so the layout has to wait
    until every picture is rendered: an ``original`` export varies in size from slide to
    slide. One shared scale keeps each picture's own aspect ratio and, between pictures,
    their relative sizes; a crop half as wide stays half as wide on its slide.
    """
    widest = max(width for _, _, (width, _) in placements)
    tallest = max(height for _, _, (_, height) in placements)
    scale = min(PICTURE_LIMIT[0] / widest, PICTURE_LIMIT[1] / tallest)
    footer = Inches(0.35) if timestamp else 0
    deck.slide_width = max(Inches(2.5 if timestamp else 1), round(widest * scale))
    deck.slide_height = max(Inches(1), round(tallest * scale) + footer)
    for picture, caption, (width, height) in placements:
        picture.width = round(width * scale)
        picture.height = round(height * scale)
        picture.left = (deck.slide_width - picture.width) // 2
        picture.top = (deck.slide_height - footer - picture.height) // 2
        if caption is not None:
            caption.left = Inches(0.15)
            caption.top = deck.slide_height - footer
            caption.width = deck.slide_width - Inches(0.3)
            caption.height = footer


def export_markers(
    info,
    markers,
    settings,
    destination,
    kind,
    cancel=None,
    progress=lambda value: None,
    overwrite=False,
):
    settings.validate()
    markers = sorted(
        [marker for marker in markers if marker.included], key=lambda marker: marker.frame.seconds
    )
    if not markers:
        raise ValueError("No markers selected for export")
    if kind not in ("jpg", "pptx"):
        raise ValueError("Unsupported output format")
    destination = Path(destination).resolve()
    if destination == info.path:
        raise ValueError("Cannot overwrite the source video")
    if kind == "pptx" and destination.exists() and not overwrite:
        raise FileExistsError(destination)
    parent = destination if kind == "jpg" else destination.parent
    with tempfile.TemporaryDirectory(prefix=".frame-export-", dir=parent) as temporary:
        staging = Path(temporary)
        deck = Presentation() if kind == "pptx" else None
        placements = []
        with FrameReader(info, cache_size=1) as reader:
            for index, marker in enumerate(markers):
                check_cancel(cancel)
                content = render_marker(reader, marker, settings, cancel)
                if deck is None:
                    name = f"{index + 1:04}_{timecode(marker.frame.seconds).replace(':', '-')}.jpg"
                    (staging / name).write_bytes(content)
                else:
                    placements.append(add_slide(deck, content, marker, settings.timestamp))
                progress(int((index + 1) / len(markers) * 100))
        check_cancel(cancel)
        if deck is not None:
            layout_deck(deck, placements, settings.timestamp)
            output = staging / "export.pptx"
            deck.save(str(output))
            check_cancel(cancel)
            if destination.exists() and not overwrite:
                raise FileExistsError(destination)
            os.replace(output, destination)
            return destination
        name = f"frames_{datetime.now().astimezone():%Y%m%d_%H%M%S}_{uuid4().hex[:6]}"
        output = destination / name
        os.rename(staging, output)
        return output
